from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image
BASE=Path('output/design-update');SRC=Path(r'D:/unity/资源/自制资源/Between Two Poles/两极之间 两格桥修订版(1).docx')
OUT=Path(r'D:/unity/资源/自制资源/Between Two Poles/两极之间 最新规则图解策划案 2026年10月08日 修订2.docx')
doc=Document(SRC)
for node in list(doc._element.body):
 if node.tag!=qn('w:sectPr'):doc._element.body.remove(node)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.page_break_before=False
 st.paragraph_format.keep_with_next=name in ['Title','Subtitle','Heading 1','Heading 2']
 st.font.name='Microsoft YaHei';st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
doc.styles['Caption'].paragraph_format.keep_with_next=False
doc.core_properties.title='两极之间 最新规则图解策划案'
doc.core_properties.subject='磁铁组合结果 推动与跨岛规则 关卡制作基准'
doc.core_properties.author='两极之间项目组';doc.core_properties.comments=''

def p(t,style=None):return doc.add_paragraph(t,style)
def h(t,new=True):
 a=p(t,'Heading 1');a.paragraph_format.page_break_before=new;return a
def sub(t):return p(t,'Heading 2')
def fig(name,caption,width=17):
 a=doc.add_paragraph();a.paragraph_format.space_after=Pt(3);a.paragraph_format.keep_with_next=True
 shape=a.add_run().add_picture(str(BASE/'figures'/(name+'.png')),width=Cm(width))
 shape._inline.docPr.set('descr',caption)
 c=p(caption,'Caption');c.paragraph_format.space_after=Pt(7)
def table(head,rows,widths):
 t=doc.add_table(rows=1,cols=len(head));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Cm(w)
 for i,txt in enumerate(head):t.rows[0].cells[i].text=txt
 for row in rows:
  cs=t.add_row().cells
  for i,txt in enumerate(row):cs[i].text=txt
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();ns=OxmlElement('w:cantSplit');pr.append(ns)
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');b.append(x)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for edge,v in [('top','85'),('bottom','85'),('left','110'),('right','110')]:
    x=OxmlElement('w:'+edge);x.set(qn('w:w'),v);x.set(qn('w:type'),'dxa');mar.append(x)
   pr.append(mar)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DFEAF1');pr.append(sh)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(0);par.paragraph_format.space_before=Pt(0);par.paragraph_format.line_spacing=1.1
    for run in par.runs:run.font.size=Pt(10);run.bold=ri==0
 return t

p('两极之间','Title');p('最新规则图解策划案','Subtitle')
p('程序与美术共用　｜　2026 年 10 月 8 日　修订 2')
sub('1  玩家在玩什么')
p('太空浮岛上的格子解谜游戏。玩家通过站位、推动、翻面和组合，把磁铁变成通路或升降装置，前往新的岛屿与高处。角色没有主动跳跃能力；同一组材料会因姿态、推动方向和接收格支撑而产生不同结果。')
p('核心循环：观察地形与磁铁 → 选择站位和推动方向 → 形成结构 → 利用通路、旋转或磁流继续前进。组合不是背配方菜单，而是理解同一套空间规则。')
fig('gallery','图 1  六种材料合成结果。图中红色为 N 极、蓝色为 S 极，红蓝互换时规则相同。')
sub('本版阅读方式')
p('先看第 2 节的基本结算；第 3 至 7 节按组合展示“条件与最终形态”；第 8 至 9 节说明通路和空间装置；第 10 节用于关卡编排与制作验收。配图展示结构关系，不按图中透视长度作为建模尺寸。')
p('本版保留两种基础材料：条形磁铁与 U 型磁铁。圆环是可移动的组合体；宽条、十字、升降台和两格桥固定在组合接收格。双 U 型与十字组成的两层升降装置属于空间摆放关系。')

h('2  格子与推动如何结算')
p('磁铁整块具有一种极性，独立物件以所在格中心定位。玩家从相邻格推动，按上下左右逐格行动；相邻静置不会自动吸附，推向占用格时才判断极性、类型、姿态与支撑。')
sub('先判断组合 再播放对应动作')
p('组合使用本次推动前的姿态，不先把推入方按普通翻滚转一遍再套配方。接收方是目标格原有的磁铁，它决定成品锚点；姿态、支撑与推动方向共同决定具体结果。动画不能改变已经判定的组合。')
table(['物件','普通推动后的状态'],[
 ('单根条形磁铁','向前一格，按推动方向翻滚 90°；可能竖立，也可能仍平躺。'),
 ('单个 U 型磁铁','向前一格，翻面 180°，最终仍平躺。没有单体竖立、侧立或自动再滚一格。'),
 ('组合圆环','平躺可推立；竖立时沿轮缘方向连续滚动，沿轮轴方向推倒。'),
 ('其他组合体','不整体平移。十字只原地旋转；半成品可补料，两格桥与升降台保持固定。')],[3.6,13.4])
sub('同极排斥与受阻反冲')
table(['接触条件','结果'],[
 ('同极且远端有支撑，后方可进入','两块各向前翻滚一格，不合成；远端允许被推到边缘外第一格。'),
 ('同极且远端有支撑，后方是石头','触发反冲，近端磁铁撞击玩家并一同向后飞出；远端留在原处。'),
 ('同极远端已经悬在边缘外','整次推动被阻挡，不继续排斥，也不因远端后方石头而反冲。'),
 ('异极且满足组合条件','在接收格生成相应组合；若姿态、占格或高度不允许，则保留原状态。')],[5.5,11.5])
p('反冲有可落地的目标岛时，磁铁落在外侧、玩家落在内侧；配置了章节出口时接续跨章节飞行。没有落地岛且不是章节出口时，漂浮 5 秒后重开当前小岛。不同高度的磁铁不能直接接触组合。')
p('操作保留撤销与重置：Z 撤销一步，R 重置当前小岛。普通行走不会跨越无支撑的空格；边缘失衡表现不提供跳跃或额外位移。')

h('3  两根平躺条形磁铁')
p('共同前提：两根条形异极，其中一根被推向另一根。这里的平行、垂直指平躺长轴之间的夹角；反向 180°也算平行。有支撑包括同高陆地与有效桥面，无支撑指边缘外悬空接收格。')
fig('flat-bars','图 2  左侧红条为推入方，中间蓝条为接收方。只改变长轴关系与接收格支撑，就会改变组合结果。')
sub('宽条是一格长的并排结构')
p('两根沿短边并排加宽，长度不增加，不是首尾接长。两条都平躺且垂直时，有支撑则推入方对齐接收方长轴；没有支撑则保持直角，生成十字。')
sub('锚点与上下关系')
p('成品锚点留在接收格。宽条长轴跟随接收方，红蓝保留各自颜色；侧向接入时，推入方保留来向一侧，避免组合后无故换边。悬空平躺十字中，接收方下沉一根条形厚度，推入方在上。')
p('宽条形成不等于通路成立。只有长边两端接通、且沿长轴进入时，才能跨过缺口；图解与判定见第 8 节。十字不作为普通跨岛桥。')

h('4  条形磁铁有竖直状态时')
p('先区分谁被推入、谁是接收方，再看平躺条的长轴是否沿推动方向。下图中，左红为推入方，右蓝为接收方；“其余”情况的边界列在图后。')
fig('standing-bars','图 3  两条都竖直形成宽条；一竖一躺的轴向特例形成宽条，其余按十字处理。',16.4)
table(['一竖一躺的具体条件','最终结果'],[
 ('竖直推入方 → 平躺接收方，平躺长轴沿推向','宽条；有无支撑都成立。'),
 ('平躺推入方 → 竖直接收方，平躺长轴沿推向','宽条；陆地与边缘外都成立，竖条倒下并排。'),
 ('一竖一躺且平躺长轴横向于推动方向','十字；原先竖直的条倒在上面。')],[12,5])
p('宽条中的竖条沿推动方向倒下，先错开并排位置再合拢，避免互相穿模。十字保留“原平躺者在下、原竖直者在上”。两条都竖直时一起倒下形成宽条。')

h('5  U 型参与的组合结果')
p('单个 U 型的最终状态始终平躺。以下配方按推动前姿态判断；红蓝、推入方与接收方可以互换，组合种类不变。成品位置仍锚定接收格。')
fig('u-combos','图 4  U 型与长条得到半成品或升降台；两个异极 U 型得到平躺圆环。')
sub('长条与 U 型')
p('两块都平躺，一定组成两格桥半成品，与向左、向右、向前、向后推无关；不因普通翻滚会把长条转起而变成升降台。长条竖直、U 型平躺，则组成单层磁流升降台，也不因推入方不同而变化。')
sub('单层磁流升降台')
p('U 型保持平躺作底座，长条在上方竖直，合成后持续产生磁流。玩家从同层入口进入柱体旁的磁流，升到 H；从高处入口也可下降。出口必须与相邻地面高度衔接，不穿过长条实体。H 为一个格子的高度单位。')
sub('两个 U 型')
p('两块异极 U 型在同格闭合成圆环，不并排加厚。初始成品平躺，红蓝各占半环。只做必要的转向与翻面；接收方不无故绕到另一侧，两个半环不得互相穿过。圆环后续推动规则见第 7 节。')

h('6  两格桥的半成品与补齐')
p('材料为一个 U 型与两根条形，三块均平躺。两根条形彼此同极，并且都与 U 型异极。先形成半成品，再用第二根平躺条形补齐 U 型的另一接口。')
fig('two-cell','图 5  两次连接得到完整两格桥。格子底色只表示占格，不代表生成实体盖板；金色细线表示接口磁流。')
sub('朝向和占格')
p('第一步无论是把条形推向 U 型，还是把 U 型推向条形，半成品都锚定接收格。延伸方向取组合前 U 型的开口方向，而不是直接取玩家推动方向。半成品和完整桥都从锚点沿该方向占两格。')
sub('第二根材料的接入')
p('第二根必须是与 U 型底座异极的平躺条形，按接口配方结算；它与第一根侧梁同极，不触发普通排斥。第一根连接永久保留，补齐后仍固定在原锚点，不可拆开或整座推动。')
sub('部署和通行')
p('普通同高陆地不阻止部署，也允许跨缺口；第二格不能覆盖石头、不同高度地面、其他占位材料或玩家。缺一侧梁时不通行；补齐后沿当前侧梁跨越，不能横穿中央区域。')
p('当前完整桥有蓝色能量面表现，但这不等于允许横穿侧梁之间。两格桥与一格宽条的跨度不同，不通过拉长磁铁、移动岸边或瞬移角色来补距离。')

h('7  圆环的组合与移动')
p('两个 U 型异极闭合后，圆环是可推动的组合体。它保留中空结构，实体环边与圆心必须区分。单个 U 型的翻面规则与圆环的竖立滚动规则分开处理。')
fig('movement','图 6  上排是单个 U 型的连续翻面；下排是圆环的平躺、竖立与推倒状态。箭头表示状态转换，具体推向取决于接触面。')
table(['圆环当前状态与受力方向','最终行为'],[
 ('平躺圆环被推动','前进一格并竖立。'),
 ('竖立圆环，沿轮缘可滚动方向推动','连续滚动，直到障碍前或边缘外第一个无支撑格停住。'),
 ('竖立圆环，沿轮轴方向推动','前进一格并倒下，恢复平躺。'),
 ('已停在边缘外，没有有效支撑','玩家不能从岸边再直接推它前进。')],[6.8,10.2])
sub('U 型闭合动画只对齐需要改变的部分')
p('沿开口方向接近时，只翻转开口不匹配的半环；左右并排且同向时，两半环相反方向转动后闭合。接收方可做必要让位，但不能让推入方绕大圈换到另一边；红蓝相对位置应与接近关系一致。')
p('圆环可沿有效桥面继续滚动或推动至另一岛。平躺环边可作为支撑，圆心不是地板；竖立后不再提供平躺通路。圆环不参与任意追加材料的连锁合成。')

h('8  边缘组合与跨岛通路')
p('单根条形、单个 U 型和圆环都能推出边缘，在第一个无支撑格保持原高度悬停。玩家停在有支撑的位置；悬空磁铁不能再被岸边玩家直接多推一格，但可接收从岸上推来的异极材料，按组合规则结算。')
fig('passage','图 7  宽条桥只沿长轴两端通行。横着放、断开一端、连接方向或高度不匹配，都不能成为这条跨岛路径。',16.3)
p('宽条长边两端可直接连接岸边，或经同向、同高的有效桥段连续接岸。接触侧面不算接通。岛上宽条仍是有高度的固定物件，不能从地面直接踩上；已有同高顶面通行另按高度判断。')
fig('transport','图 8  通路有效时，可推动的磁铁可以推上桥、继续过桥并到达对岸。桥本身不作为这次推动的合成或排斥对象。',16.3)
p('桥上材料与桥面分层占位。上桥、过桥、下桥都要衔接支撑高度；人物跟随进入腾出的可站立位置。横向宽条、未完成两格桥、十字、石头和不匹配高差不能冒充可通过的桥面。')

h('9  十字旋转与两层磁流')
p('十字固定在接收格，不可直接走过，不可被当成宽条桥。未激活磁流时，从侧面推动十字触发原地顺时针旋转 90°；旋转不交换上下层和颜色。')
fig('cross-rotor','图 9  上图是十字拨动竖直单条；下图是双 U 型相向围绕十字的空间装置。两层装置不是三者推入同格。')
sub('旋转影响哪些物件')
p('十字只拨动相邻同层、竖直的单根条形，使其沿远离十字中心的方向移动一格，保持竖直。平躺条形低于作用高度，不被拨动；U 型和组合体不能被这样拨走。相关邻格有重型物件，或竖条目标格无同高地面、被占用、被石头或主角挡住时，十字受阻回位，不触发连锁推动。')
sub('三个位置形成两层装置')
p('中央十字，两侧相邻位置各放一个平躺 U 型；两 U 型异极、同高、开口都朝向十字。左右或前后布置均成立。三者保持独立，总共使用四块基础磁铁，不合成大圆环或新平台。')
p('条件成立时，十字持续旋转，产生升高 2H 的磁流；玩家从旁侧入口进入，可往返升降。两侧同极、开口背离、位置或高度不对时不激活。装置未激活时，十字回到普通侧推旋转规则。')

h('10  制作验收与四章规划')
sub('程序与美术共同验收')
table(['检查项','应得到的结果'],[
 ('姿态与方向','按推动前姿态合成；四向推动、红蓝互换和反向长轴都符合第 3 至 5 节。'),
 ('接收格与材料关系','成品锚定接收格；只消耗正确材料；普通游戏不拆解，撤销可恢复原状态。'),
 ('条形特殊情况','平躺垂直在陆地成宽条、边缘外成十字；一竖一躺必须区分接收方与轴向特例。'),
 ('桥的朝向与支撑','宽条必须沿长轴两端接通；两格桥随原 U 型开口延伸；半成品不通行。'),
 ('移动与边缘','单 U 型每推一次只翻一格；圆环能推立、滚动、推倒；无支撑时不能再次直接推动。'),
 ('组合动画','保持原颜色与来向关系；只做必要的翻转和让位；不穿模、不绕大圈、不在换模型时跳位。'),
 ('跨岛与高差','磁铁能经有效桥面运输；同级悬空材料阻挡继续外推；磁流入口与出口可达。'),
 ('撤销与重置','组合前后、过桥、圆环滚动和边缘悬停均可正确撤销；小岛重置恢复相应材料。')],[3.6,13.4])
sub('四章功能规划')
table(['章节','介绍内容与关卡重点'],[
 ('第一章  基础操作与反冲','只用条形入门，学习推动、姿态、同极排斥和异极组合；受阻反冲用于跨岛及章节衔接。原目标时长约 10 分钟。'),
 ('第二章  宽条与十字','比较陆地与边缘外的条形结果，学习接收方、桥两端连接、十字旋转与竖条拨动。'),
 ('第三章  U 型与跨岛运输','引入 U 型翻面、圆环移动、半成品补齐、两格桥与沿桥运输材料。圆环作为可移动组合体单独说明。'),
 ('第四章  磁流与高度','单层升降台与高低平台，后段加入双 U 型围绕十字的两层磁流，组合高度与空间站位谜题。')],[4.6,12.4])
p('美术保持斜俯视、深色太空、低模浮岛与清楚的格子边缘。优先让玩家看清长轴、U 型开口、桥端、圆心空洞和上下层；光效不能遮住角色、极性与落脚点。')
p('版本基准：本案替换旧版中“先翻滚后判配方”“单 U 型竖立滚动”“所有组合体不可移动”和“圆环用途未定义”等规则。剧情与分镜另见独立文档；四章条目为关卡编排规划。')

# Remove stale direct theme coloring from source furniture; retain page fields.
for section in doc.sections:
 for footer in [section.footer,section.first_page_footer,section.even_page_footer]:
  for par in footer.paragraphs:
   for run in par.runs:run.font.color.rgb=RGBColor(0,0,0)
settings=doc.settings.element
uf=OxmlElement('w:updateFields');uf.set(qn('w:val'),'true');settings.append(uf)
doc.save(OUT)
print(str(OUT))
